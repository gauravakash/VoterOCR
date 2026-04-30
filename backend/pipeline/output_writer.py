"""Write a BoothOutput to JSON + multi-sheet Excel with analytics on Sheet 2.

Excel layout:
  1  All Voters
  2  📊 Analytics (NEW PRIMARY VIEW — embeds religion pie + caste/surname bar charts)
  3  Religion Summary
  4  Hindu Castes
  5  Caste Categories
  6  Top 10 Surnames
  7  Rare Surnames
  8  Unknown Caste (Review)
  9  Validation
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill

from .. import config
from ..schemas.voter import BoothOutput
from .analytics import compute_booth_analytics, get_unknown_caste_voters


VOTER_COLUMNS = [
    "serial_no", "voter_id", "name", "relation_type", "relative_name",
    "house_no", "age", "gender",
    "religion", "religion_confidence",
    "caste", "caste_category", "caste_confidence", "surname_matched",
    "booth_number", "booth_name", "page_no", "extraction_model",
]

ANALYTICS_SHEET_NAME = "📊 Analytics"


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_SECTION_FONT = Font(bold=True, size=13, color="1F2937")


def _style_header_row(ws, row_num: int = 1) -> None:
    for cell in ws[row_num]:
        if cell.value is None:
            continue
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize_columns(ws, max_width: int = 50) -> None:
    for col in ws.columns:
        cells = [c for c in col if c.value is not None]
        if not cells:
            continue
        length = max(len(str(c.value)) for c in cells)
        ws.column_dimensions[cells[0].column_letter].width = min(length + 2, max_width)


# ---------------------------------------------------------------------------
# Analytics sheet (Sheet 2)
# ---------------------------------------------------------------------------

def _write_analytics_sheet(wb, analytics: dict, metadata: dict, validation: dict, routing: dict | None = None) -> None:
    ws = wb.create_sheet(ANALYTICS_SHEET_NAME, 1)

    ws["A1"] = "BOOTH ANALYTICS REPORT"
    ws["A1"].font = Font(bold=True, size=16, color="1F2937")
    ws.merge_cells("A1:D1")

    header_rows = [
        ("Booth", metadata.get("booth_name") or "N/A"),
        ("Constituency", metadata.get("constituency") or "N/A"),
        ("Booth Number", metadata.get("booth_number") or "N/A"),
        ("Total Voters", analytics.get("total_voters", 0)),
        ("Validation", validation.get("status", "N/A")),
    ]
    for i, (label, value) in enumerate(header_rows, start=2):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)

    row = 8

    # ===== Religion =====
    ws.cell(row=row, column=1, value="RELIGION DISTRIBUTION").font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    ws.cell(row=row, column=1, value="Religion")
    ws.cell(row=row, column=2, value="Count")
    ws.cell(row=row, column=3, value="Percentage")
    _style_header_row(ws, row)
    religion_header_row = row
    row += 1
    religion_start = row
    for k, v in sorted(analytics.get("religion", {}).items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        ws.cell(row=row, column=3, value=f"{analytics.get('religion_percentages', {}).get(k, 0)}%")
        row += 1
    religion_end = row - 1

    if religion_end >= religion_start:
        pie = PieChart()
        labels = Reference(ws, min_col=1, min_row=religion_start, max_row=religion_end)
        data = Reference(ws, min_col=2, min_row=religion_header_row, max_row=religion_end)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = "Religion Distribution"
        pie.height = 9
        pie.width = 13
        pie.dataLabels = DataLabelList(showPercent=True)
        ws.add_chart(pie, f"E{religion_header_row - 1}")

    row += 2

    # ===== Caste Categories =====
    ws.cell(row=row, column=1, value="HINDU CASTE CATEGORIES").font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    ws.cell(row=row, column=1, value="Category")
    ws.cell(row=row, column=2, value="Count")
    _style_header_row(ws, row)
    cat_header_row = row
    row += 1
    cat_start = row
    cats = analytics.get("hindu_caste_categories", {})
    for k, v in sorted(cats.items(), key=lambda x: -x[1]):
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        row += 1
    cat_end = row - 1

    if cats and cat_end >= cat_start:
        bar = BarChart()
        labels = Reference(ws, min_col=1, min_row=cat_start, max_row=cat_end)
        data = Reference(ws, min_col=2, min_row=cat_header_row, max_row=cat_end)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(labels)
        bar.title = "Hindu Caste Categories (UC/OBC/SC)"
        bar.height = 9
        bar.width = 13
        ws.add_chart(bar, f"E{cat_header_row - 1}")

    row += 2

    # ===== Top 10 Surnames =====
    ws.cell(row=row, column=1, value="TOP 10 SURNAMES").font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    ws.cell(row=row, column=1, value="Surname")
    ws.cell(row=row, column=2, value="Count")
    _style_header_row(ws, row)
    sn_header_row = row
    row += 1
    sn_start = row
    for item in analytics.get("top_10_surnames", []):
        ws.cell(row=row, column=1, value=item.get("surname"))
        ws.cell(row=row, column=2, value=item.get("count"))
        row += 1
    sn_end = row - 1

    if sn_end >= sn_start:
        bar = BarChart()
        labels = Reference(ws, min_col=1, min_row=sn_start, max_row=sn_end)
        data = Reference(ws, min_col=2, min_row=sn_header_row, max_row=sn_end)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(labels)
        bar.title = "Top 10 Surnames"
        bar.height = 9
        bar.width = 13
        ws.add_chart(bar, f"E{sn_header_row - 1}")

    row += 2

    # ===== Hindu Specific Castes (top 15) =====
    ws.cell(row=row, column=1, value="HINDU SPECIFIC CASTES (TOP 15)").font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    ws.cell(row=row, column=1, value="Caste")
    ws.cell(row=row, column=2, value="Count")
    _style_header_row(ws, row)
    hc_header_row = row
    row += 1
    hc_start = row
    castes = analytics.get("hindu_castes", {})
    for k, v in sorted(castes.items(), key=lambda x: -x[1])[:15]:
        ws.cell(row=row, column=1, value=k or "Unknown")
        ws.cell(row=row, column=2, value=v)
        row += 1
    hc_end = row - 1

    if castes and hc_end >= hc_start:
        bar = BarChart()
        labels = Reference(ws, min_col=1, min_row=hc_start, max_row=hc_end)
        data = Reference(ws, min_col=2, min_row=hc_header_row, max_row=hc_end)
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(labels)
        bar.title = "Hindu Caste Distribution"
        bar.height = 11
        bar.width = 14
        ws.add_chart(bar, f"E{hc_header_row - 1}")

    row += 2

    # ===== Key Metrics =====
    ws.cell(row=row, column=1, value="KEY METRICS").font = _SECTION_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1
    metrics = [
        ("Total Voters", analytics.get("total_voters", 0)),
        ("Hindu", analytics.get("religion", {}).get("Hindu", 0)),
        ("Muslim", analytics.get("religion", {}).get("Muslim", 0)),
        ("Sikh", analytics.get("religion", {}).get("Sikh", 0)),
        ("Unknown Religion", analytics.get("religion", {}).get("Unknown", 0)),
        ("Hindus w/ Unknown Caste", analytics.get("unknown_caste_count", 0)),
        ("Male", analytics.get("gender", {}).get("Male", 0)),
        ("Female", analytics.get("gender", {}).get("Female", 0)),
    ]
    for label, val in metrics:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=val)
        row += 1

    # ===== Cost Breakdown (actual, from cost_tracker) =====
    if routing:
        row += 2
        ws.cell(row=row, column=1, value="COST BREAKDOWN (ACTUAL)").font = _SECTION_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        row += 1
        cost_rows = [
            ("Total Cost (INR)", f"₹{routing.get('cost_inr', 0):.2f}"),
            ("Total Cost (USD)", f"${routing.get('cost_usd', 0):.4f}"),
            ("Input Tokens",  f"{routing.get('input_tokens', 0):,}"),
            ("Output Tokens", f"{routing.get('output_tokens', 0):,}"),
            ("Cached Tokens", f"{routing.get('cached_tokens', 0):,}"),
            ("Flash Pages",         routing.get("flash_pages", 0)),
            ("Pro Fallback Pages",  routing.get("pro_fallback_pages", 0)),
            ("Failed Pages",        routing.get("failed_pages", 0)),
            ("Fallback Rate (%)",   routing.get("fallback_rate_pct", 0)),
        ]
        for label, val in cost_rows:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=val)
            row += 1

    _autosize_columns(ws, max_width=40)


# ---------------------------------------------------------------------------
# Top-level write
# ---------------------------------------------------------------------------

def save_outputs(
    output: BoothOutput,
    output_dir: Path = config.OUTPUT_DIR,
) -> tuple[Path, Path, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(output.pdf_source).stem

    json_path = output_dir / f"{stem}_voters.json"
    xlsx_path = output_dir / f"{stem}_voters.xlsx"
    summary_path = output_dir / f"{stem}_summary.json"

    voters_data = [v.model_dump() for v in output.voters]
    analytics = compute_booth_analytics(voters_data)

    payload = output.model_dump(mode="json")
    payload["analytics"] = analytics
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2, default=str)

    summary = {
        "pdf_source": output.pdf_source,
        "total_pages": output.total_pages,
        "metadata": payload["metadata"],
        "validation": payload["validation"],
        "analytics": analytics,
        "llm_provider": output.llm_provider,
        "llm_model": output.llm_model,
        "processed_at": payload["processed_at"],
    }
    with open(summary_path, "w", encoding="utf-8") as fh:
        json.dump(summary, fh, ensure_ascii=False, indent=2, default=str)

    _write_excel(xlsx_path, output, voters_data, analytics)

    return json_path, xlsx_path, summary_path


def _write_excel(path: Path, output: BoothOutput, voters_data: list[dict], analytics: dict) -> None:
    df = pd.DataFrame(voters_data)
    for c in VOTER_COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[VOTER_COLUMNS] if not df.empty else pd.DataFrame(columns=VOTER_COLUMNS)

    val = output.validation
    meta = output.metadata

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Sheet 1
        df.to_excel(writer, sheet_name="All Voters", index=False)
        _style_header_row(writer.book["All Voters"])
        _autosize_columns(writer.book["All Voters"])

        # Sheets 3+ (Sheet 2 is inserted at end after charts can reference)
        religion_rows = [
            {
                "Religion": k,
                "Count": v,
                "Percentage": analytics["religion_percentages"].get(k, 0),
            }
            for k, v in sorted(analytics["religion"].items(), key=lambda x: -x[1])
        ]
        if religion_rows:
            pd.DataFrame(religion_rows).to_excel(writer, sheet_name="Religion Summary", index=False)
            _style_header_row(writer.book["Religion Summary"])

        if analytics.get("hindu_castes"):
            pd.DataFrame([
                {"Caste": k or "Unknown", "Count": v}
                for k, v in sorted(analytics["hindu_castes"].items(), key=lambda x: -x[1])
            ]).to_excel(writer, sheet_name="Hindu Castes", index=False)
            _style_header_row(writer.book["Hindu Castes"])

        if analytics.get("hindu_caste_categories"):
            pd.DataFrame([
                {"Category": k, "Count": v}
                for k, v in sorted(analytics["hindu_caste_categories"].items(), key=lambda x: -x[1])
            ]).to_excel(writer, sheet_name="Caste Categories", index=False)
            _style_header_row(writer.book["Caste Categories"])

        if analytics.get("top_10_surnames"):
            pd.DataFrame(analytics["top_10_surnames"]).to_excel(
                writer, sheet_name="Top 10 Surnames", index=False
            )
            _style_header_row(writer.book["Top 10 Surnames"])

        if analytics.get("rare_surnames"):
            pd.DataFrame(analytics["rare_surnames"]).to_excel(
                writer, sheet_name="Rare Surnames", index=False
            )
            _style_header_row(writer.book["Rare Surnames"])

        unknown = get_unknown_caste_voters(voters_data)
        if unknown:
            unk_df = pd.DataFrame(unknown)
            for c in VOTER_COLUMNS:
                if c not in unk_df.columns:
                    unk_df[c] = None
            unk_df[VOTER_COLUMNS].to_excel(
                writer, sheet_name="Unknown Caste", index=False
            )
            _style_header_row(writer.book["Unknown Caste"])

        # Validation sheet (last)
        val_rows = [
            ("PDF", output.pdf_source),
            ("Constituency", meta.constituency or ""),
            ("Booth Number", meta.booth_number or ""),
            ("Booth Name", meta.booth_name or ""),
            ("Pin Code", meta.pin_code or ""),
            ("", ""),
            ("Expected Voters (page 1)", val.expected_total),
            ("Extracted Voters", val.extracted_total),
            ("Difference", val.difference),
            ("Match", val.match),
            ("Status", val.status),
            ("", ""),
            ("Expected Male", val.expected_male),
            ("Extracted Male", val.extracted_male),
            ("Expected Female", val.expected_female),
            ("Extracted Female", val.extracted_female),
            ("", ""),
            ("Missing Serial Numbers (first 50)", str(val.missing_serial_numbers)),
            ("Missing Serial Count", val.missing_serial_count),
            ("Duplicate Serials", str(val.duplicate_serials)),
            ("Duplicate Voter IDs", str(val.duplicate_voter_ids)),
            ("Duplicates Removed (dedup)", val.duplicates_removed),
            ("", ""),
            ("LLM Provider", output.llm_provider),
            ("LLM Model", output.llm_model),
            ("Processed At", str(output.processed_at)),
        ]
        if output.routing_stats:
            rs = output.routing_stats
            val_rows.extend([
                ("", ""),
                ("Flash Pages", rs.flash_pages),
                ("Pro Fallback Pages", rs.pro_fallback_pages),
                ("Failed Pages", rs.failed_pages),
                ("Fallback Rate (%)", rs.fallback_rate_pct),
                ("", ""),
                ("Cost (INR, actual)", f"₹{rs.cost_inr:.2f}"),
                ("Cost (USD, actual)", f"${rs.cost_usd:.4f}"),
                ("Input Tokens", f"{rs.input_tokens:,}"),
                ("Output Tokens", f"{rs.output_tokens:,}"),
                ("Cached Tokens", f"{rs.cached_tokens:,}"),
            ])
        pd.DataFrame(val_rows, columns=["Metric", "Value"]).to_excel(
            writer, sheet_name="Validation", index=False
        )
        _style_header_row(writer.book["Validation"])

        # Sheet 2 (Analytics) — inserted at position 1 so it appears 2nd in the tab bar
        routing_dict = output.routing_stats.model_dump() if output.routing_stats else None
        _write_analytics_sheet(writer.book, analytics, meta.model_dump(), val.model_dump(), routing_dict)
